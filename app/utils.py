"""通用工具函数。"""
import json


def safe_json_dumps(data) -> str:
    """序列化为 JSON 并转义尖括号，防止 </script> 突破 script 块造成存储型 XSS。

    模板中仍需配合 |safe 使用（内容已是合法 JSON，转义后不影响 JSON.parse）。
    """
    return (
        json.dumps(data, ensure_ascii=False)
        .replace('<', '\\u003c')
        .replace('>', '\\u003e')
        .replace('&', '\\u0026')
    )


def parse_float_or(value: str | None, default: float = 0.0) -> float:
    """表单数字框安全解析：空字符串/非法值回退为 default，避免 Pydantic 422。"""
    if value is None:
        return default
    text = str(value).strip()
    if not text:
        return default
    try:
        return float(text)
    except (TypeError, ValueError):
        return default


def parse_optional_int(value: str | None) -> int | None:
    """表单整数框安全解析：空字符串/非法值返回 None。"""
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    try:
        return int(float(text))
    except (TypeError, ValueError):
        return None


def safe_excel_cell(value):
    """Excel 公式注入防护：以 = + - @ 开头的字符串前加单引号，强制按文本处理。"""
    if isinstance(value, str) and value[:1] in ('=', '+', '-', '@'):
        return "'" + value
    return value


def xlsx_response(filename: str, sheets: list):
    """把 [(sheet名, 表头, 列宽, 行数据), ...] 打包成 xlsx 下载响应。

    文本单元格统一过 safe_excel_cell 防公式注入；表头蓝底加粗，与既有导出口径一致。
    """
    import io
    from urllib.parse import quote

    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill('solid', fgColor='D9E1F2')
    header_font = Font(bold=True)
    for title, headers, widths, rows in sheets:
        ws = wb.create_sheet(title=title)
        for col, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
        for r_idx, row in enumerate(rows, start=2):
            for c_idx, value in enumerate(row, start=1):
                ws.cell(row=r_idx, column=c_idx, value=safe_excel_cell(value))
        for col, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(col)].width = width
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f"attachment; filename*=UTF-8''{quote(filename)}"},
    )
