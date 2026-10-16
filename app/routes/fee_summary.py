import io
from datetime import date
from urllib.parse import quote

from fastapi import APIRouter, Depends, Request
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy import and_, case, extract, func
from sqlalchemy.orm import Session

from app.database import get_db
from app.utils import safe_excel_cell
from app.models import Company, FeeRecord, PipelineEntry
from app.paths import TEMPLATES_DIR

router = APIRouter(prefix='/fee-summary', tags=['fee_summary'])
templates = Jinja2Templates(directory=str(TEMPLATES_DIR))


def get_available_years(db: Session) -> list[int]:
    """年度下拉：从实收时间、应收时间里提取真实存在的年份，避免写死 2024~2041。"""
    received_years = (
        db.query(extract('year', FeeRecord.actual_received_date).label('y'))
        .filter(FeeRecord.actual_received_date.isnot(None))
        .distinct()
        .all()
    )
    planned_years = (
        db.query(extract('year', FeeRecord.planned_receivable_date).label('y'))
        .filter(FeeRecord.planned_receivable_date.isnot(None))
        .distinct()
        .all()
    )
    years = sorted({int(row.y) for row in received_years + planned_years if row.y is not None}, reverse=True)
    return years or [date.today().year]


def outstanding_expr():
    """未收余额表达式：含税应收 - 实收。"""
    return FeeRecord.amount_incl_tax - func.coalesce(FeeRecord.actual_received_amount, 0)


def build_received_query(db: Session, selected_year: int | None):
    return (
        db.query(
            extract('year', FeeRecord.actual_received_date).label('year'),
            Company.id.label('company_id'),
            Company.company_name.label('company_name'),
            Company.short_name.label('short_name'),
            func.count(FeeRecord.id).label('record_count'),
            func.coalesce(func.sum(FeeRecord.actual_received_amount), 0).label('total_received'),
        )
        .join(Company, Company.id == FeeRecord.company_id)
        .filter(FeeRecord.actual_received_date.isnot(None))
        .filter(FeeRecord.actual_received_amount.isnot(None))
        .filter(FeeRecord.actual_received_amount > 0)
        .filter(extract('year', FeeRecord.actual_received_date) == selected_year if selected_year else True)
    )


MAINTENANCE_FEE_TYPES = ['日常维护费', '运维费', '维护费']


def build_project_received_query(db: Session, selected_year: int | None):
    """项目汇总：按项目 + 年度分组，拆分入廊费 / 日常维护费实收。"""
    project_name = func.coalesce(FeeRecord.project_name, PipelineEntry.project_name).label('project_name')
    entry_received = func.coalesce(
        func.sum(case((FeeRecord.fee_type == '入廊费', FeeRecord.actual_received_amount), else_=0)), 0
    ).label('entry_received')
    maint_received = func.coalesce(
        func.sum(case((FeeRecord.fee_type.in_(MAINTENANCE_FEE_TYPES), FeeRecord.actual_received_amount), else_=0)), 0
    ).label('maint_received')
    return (
        db.query(
            extract('year', FeeRecord.actual_received_date).label('year'),
            project_name,
            Company.company_name.label('company_name'),
            Company.short_name.label('short_name'),
            func.count(FeeRecord.id).label('record_count'),
            entry_received,
            maint_received,
            func.coalesce(func.sum(FeeRecord.actual_received_amount), 0).label('total_received'),
        )
        .join(Company, Company.id == FeeRecord.company_id)
        .outerjoin(PipelineEntry, PipelineEntry.id == FeeRecord.pipeline_entry_id)
        .filter(FeeRecord.actual_received_date.isnot(None))
        .filter(FeeRecord.actual_received_amount.isnot(None))
        .filter(FeeRecord.actual_received_amount > 0)
        .filter(extract('year', FeeRecord.actual_received_date) == selected_year if selected_year else True)
        .group_by(
            extract('year', FeeRecord.actual_received_date),
            project_name,
            Company.id,
            Company.company_name,
            Company.short_name,
        )
        .order_by(
            extract('year', FeeRecord.actual_received_date).desc(),
            func.sum(FeeRecord.actual_received_amount).desc(),
        )
    )


def build_pending_query(db: Session, selected_year: int | None, today: date):
    """待收/逾期汇总：按应收年度 + 单位分组，只统计仍有未收余额的记录。
    无应收日期但有欠款的归入"未定"年度组，与首页待收口径保持一致。"""
    outstanding = outstanding_expr()
    overdue_amount = case(
        (and_(FeeRecord.planned_receivable_date < today, outstanding > 0), outstanding),
        else_=0,
    )
    return (
        db.query(
            extract('year', FeeRecord.planned_receivable_date).label('year'),
            Company.id.label('company_id'),
            Company.company_name.label('company_name'),
            Company.short_name.label('short_name'),
            func.count(FeeRecord.id).label('record_count'),
            func.coalesce(func.sum(FeeRecord.amount_incl_tax), 0).label('total_receivable'),
            func.coalesce(func.sum(outstanding), 0).label('total_unreceived'),
            func.coalesce(func.sum(overdue_amount), 0).label('total_overdue'),
        )
        .join(Company, Company.id == FeeRecord.company_id)
        .filter(outstanding > 0)
        .filter(extract('year', FeeRecord.planned_receivable_date) == selected_year if selected_year else True)
    )


@router.get('/', response_class=HTMLResponse)
def fee_summary(request: Request, db: Session = Depends(get_db), year: int | None = None):
    years = get_available_years(db)
    selected_year = year if year in years else None
    today = date.today()

    summary_rows = (
        build_received_query(db, selected_year)
        .group_by(
            extract('year', FeeRecord.actual_received_date),
            Company.id,
            Company.company_name,
            Company.short_name,
        )
        .order_by(
            extract('year', FeeRecord.actual_received_date).desc(),
            func.sum(FeeRecord.actual_received_amount).desc(),
        )
        .all()
    )

    pending_rows = (
        build_pending_query(db, selected_year, today)
        .group_by(
            extract('year', FeeRecord.planned_receivable_date),
            Company.id,
            Company.company_name,
            Company.short_name,
        )
        .order_by(
            extract('year', FeeRecord.planned_receivable_date).desc(),
            func.sum(outstanding_expr()).desc(),
        )
        .all()
    )

    project_rows = build_project_received_query(db, selected_year).all()

    project_entry_total = sum(float(row.entry_received or 0) for row in project_rows)
    project_maint_total = sum(float(row.maint_received or 0) for row in project_rows)
    project_total = sum(float(row.total_received or 0) for row in project_rows)

    total_received = sum(float(row.total_received or 0) for row in summary_rows)
    total_pending_receivable = sum(float(row.total_receivable or 0) for row in pending_rows)
    total_pending_unreceived = sum(float(row.total_unreceived or 0) for row in pending_rows)
    total_pending_overdue = sum(float(row.total_overdue or 0) for row in pending_rows)
    company_count = len({row.company_id for row in summary_rows})

    return templates.TemplateResponse(
        'fee_summary/list.html',
        {
            'request': request,
            'summary_rows': summary_rows,
            'pending_rows': pending_rows,
            'project_rows': project_rows,
            'project_entry_total': project_entry_total,
            'project_maint_total': project_maint_total,
            'project_total': project_total,
            'years': years,
            'selected_year': selected_year,
            'total_received': total_received,
            'total_pending_receivable': total_pending_receivable,
            'total_pending_unreceived': total_pending_unreceived,
            'total_pending_overdue': total_pending_overdue,
            'company_count': company_count,
            'title': '收费汇总',
        },
    )


@router.get('/export')
def export_summary(request: Request, db: Session = Depends(get_db), year: int | None = None):
    """导出收费汇总为 Excel（含实收汇总、项目汇总、待收/逾期汇总三个工作表）。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    years = get_available_years(db)
    selected_year = year if year in years else None
    today = date.today()

    summary_rows = (
        build_received_query(db, selected_year)
        .group_by(
            extract('year', FeeRecord.actual_received_date),
            Company.id,
            Company.company_name,
            Company.short_name,
        )
        .order_by(
            extract('year', FeeRecord.actual_received_date).desc(),
            func.sum(FeeRecord.actual_received_amount).desc(),
        )
        .all()
    )
    pending_rows = (
        build_pending_query(db, selected_year, today)
        .group_by(
            extract('year', FeeRecord.planned_receivable_date),
            Company.id,
            Company.company_name,
            Company.short_name,
        )
        .order_by(
            extract('year', FeeRecord.planned_receivable_date).desc(),
            func.sum(outstanding_expr()).desc(),
        )
        .all()
    )

    def write_sheet(ws, headers, widths, rows):
        header_fill = PatternFill('solid', fgColor='D9E1F2')
        header_font = Font(bold=True)
        for col, title in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=title)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
        for col, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(col)].width = width
        for row_idx, values in enumerate(rows, start=2):
            for col, value in enumerate(values, start=1):
                ws.cell(row=row_idx, column=col, value=safe_excel_cell(value))

    wb = Workbook()

    ws1 = wb.active
    ws1.title = '实收汇总'
    write_sheet(
        ws1,
        ['年度', '单位名称', '单位简称', '收费记录笔数', '实收总金额'],
        [10, 28, 14, 14, 16],
        [
            [int(row.year), row.company_name, row.short_name or '',
             row.record_count, float(row.total_received or 0)]
            for row in summary_rows
        ],
    )

    project_rows = build_project_received_query(db, selected_year).all()

    ws3 = wb.create_sheet('项目汇总')
    write_sheet(
        ws3,
        ['年度', '项目名称', '单位名称', '单位简称', '入廊费实收', '日常维护费实收', '实收合计', '笔数'],
        [10, 30, 28, 14, 16, 16, 16, 10],
        [
            [int(row.year), row.project_name or '', row.company_name, row.short_name or '',
             float(row.entry_received or 0), float(row.maint_received or 0),
             float(row.total_received or 0), row.record_count]
            for row in project_rows
        ],
    )

    ws2 = wb.create_sheet('待收逾期汇总')
    write_sheet(
        ws2,
        ['年度(应收)', '单位名称', '单位简称', '待收笔数', '应收总额(含税)', '未收金额', '逾期金额'],
        [12, 28, 14, 12, 16, 14, 14],
        [
            [int(row.year) if row.year else '未定', row.company_name, row.short_name or '',
             row.record_count, float(row.total_receivable or 0),
             float(row.total_unreceived or 0), float(row.total_overdue or 0)]
            for row in pending_rows
        ],
    )

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = quote('收费汇总.xlsx')
    return StreamingResponse(
        buf,
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f"attachment; filename*=UTF-8''{filename}"},
    )
