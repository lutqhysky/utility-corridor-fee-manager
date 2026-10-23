import json
from datetime import date, datetime
from fastapi import APIRouter, Depends, Form, HTTPException, Request
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session
from urllib.parse import quote
from app.database import get_db
from app.utils import safe_excel_cell, safe_json_dumps
# 引入了 FeeStandard 模型
from app.models import Company, FeeRecord, PipelineEntry, FeeStandard 
from app.paths import TEMPLATES_DIR
from app.services.reminder_service import (
    get_reminder_settings,
    get_setting,
    run_daily_digest,
    run_fee_reminders,
)

router = APIRouter(prefix='/fee-records', tags=['fee_records'])
templates = Jinja2Templates(directory=str(TEMPLATES_DIR))

# 动态逾期筛选的特殊取值（非数据库状态，而是按“应收时间<今天且有未收余额”实时判定）
OVERDUE_FILTER_VALUE = '__overdue__'


# 维护类费用类型的别名归一化（表单/历史数据用"运维费"，收费标准用"日常维护费"）
MAINTENANCE_FEE_TYPES = {'运维费', '日常维护费', '维护费'}


def normalize_fee_type(fee_type: str | None) -> str:
    return (fee_type or '').strip()


def get_standard_tax_rate(db: Session, fee_type: str | None) -> float | None:
    """按费用类型取启用中收费标准的税率：取最新生效的一条，找不到返回 None。"""
    fee_type = normalize_fee_type(fee_type)
    candidates = {fee_type}
    if fee_type in MAINTENANCE_FEE_TYPES:
        candidates |= MAINTENANCE_FEE_TYPES
    row = (
        db.query(FeeStandard)
        .filter(FeeStandard.is_active.is_(True), FeeStandard.fee_type.in_(candidates))
        .order_by(
            FeeStandard.effective_date.isnot(None).desc(),
            FeeStandard.effective_date.desc(),
            FeeStandard.id.desc(),
        )
        .first()
    )
    if row is not None and row.tax_rate is not None:
        return float(row.tax_rate)
    return None


def get_all_standard_tax_rates(db: Session) -> dict:
    """供表单回显：各费用类型的当前标准税率。"""
    return {
        '入廊费': get_standard_tax_rate(db, '入廊费'),
        '日常维护费': get_standard_tax_rate(db, '日常维护费'),
    }


def is_record_overdue(record: FeeRecord, today: date) -> bool:
    """动态判定一条收费记录是否已逾期：有应收时间、已过应收时间、仍有未收余额、且未标记已收缴。"""
    if record.payment_status == '已收缴':
        return False
    if not record.planned_receivable_date or record.planned_receivable_date >= today:
        return False
    outstanding = float(record.amount_incl_tax or 0) - float(record.actual_received_amount or 0)
    return outstanding > 0


def build_filtered_summary(records: list[FeeRecord]) -> tuple[list[dict], dict]:
    """按（单位, 项目）分组汇总实收：入廊费 / 日常维护费（含别名）/ 合计。"""
    groups: dict[tuple[str, str], dict] = {}
    for record in records:
        company = record.company
        company_label = (company.short_name or company.company_name) if company else ''
        project = (record.project_name or '').strip()
        key = (company_label, project)
        group = groups.setdefault(key, {'company': company_label, 'project': project,
                                        'entry': 0.0, 'maint': 0.0, 'total': 0.0, 'count': 0})
        amount = float(record.actual_received_amount or 0)
        group['total'] += amount
        group['count'] += 1
        fee_type = (record.fee_type or '').strip()
        if fee_type == '入廊费':
            group['entry'] += amount
        elif fee_type in MAINTENANCE_FEE_TYPES:
            group['maint'] += amount
    group_list = sorted(groups.values(), key=lambda g: g['total'], reverse=True)
    total = {
        'entry': sum(g['entry'] for g in group_list),
        'maint': sum(g['maint'] for g in group_list),
        'total': sum(g['total'] for g in group_list),
        'count': sum(g['count'] for g in group_list),
    }
    return group_list, total


def annotate_overdue(records: list[FeeRecord], today: date) -> int:
    """给记录打上 is_overdue_view 动态逾期标记，返回逾期笔数。"""
    count = 0
    for record in records:
        record.is_overdue_view = is_record_overdue(record, today)
        if record.is_overdue_view:
            count += 1
    return count


def parse_date(value: str):
    if not value:
        return None
    try:
        return datetime.strptime(value, '%Y-%m-%d').date()
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=f'日期格式错误: {value}，正确格式应为 YYYY-MM-DD') from exc


def parse_float_or(value: str | None, default: float = 0.0) -> float:
    """数字框留空（""）时用默认值，避免 422。"""
    if value is None:
        return default
    stripped = value.strip()
    if not stripped:
        return default
    try:
        return float(stripped)
    except (ValueError, TypeError):
        return default


def parse_optional_int(value: str) -> int | None:
    if value is None:
        return None
    stripped = value.strip()
    if not stripped:
        return None
    try:
        return int(stripped)
    except ValueError:
        return None


def calculate_entry_actual_fees(entry: PipelineEntry) -> tuple[float, float]:
    entry_fee_excl_tax_total = sum((item.entry_amount_excl_tax or 0) for item in entry.details)
    maintenance_fee_excl_tax_total = sum((item.maintenance_amount_excl_tax or 0) for item in entry.details)

    entry_fee_tax_rate = entry.entry_fee_tax_rate or 0
    maintenance_fee_tax_rate = entry.maintenance_fee_tax_rate or 0

    entry_fee_discount = entry.entry_fee_discount if entry.entry_fee_discount is not None else 1
    maintenance_fee_discount = entry.maintenance_fee_discount if entry.maintenance_fee_discount is not None else 1

    # 与入廊清单详情页统一口径：Decimal 分步舍入（税金→含税→打折）
    from app.services.fee_calc_service import calc_actual_fee
    entry_actual_fee = calc_actual_fee(entry_fee_excl_tax_total, entry_fee_tax_rate, entry_fee_discount)
    maintenance_actual_fee = calc_actual_fee(maintenance_fee_excl_tax_total, maintenance_fee_tax_rate, maintenance_fee_discount)
    return float(entry_actual_fee), float(maintenance_actual_fee)


def calc_excl_tax_from_incl_tax(amount_incl_tax: float, tax_rate: float) -> tuple[float, float]:
    """含税反算不含税/税金：Decimal + 四舍五入，与 fee_calc_service 保持一致。"""
    from decimal import Decimal, ROUND_HALF_UP
    rate = Decimal(str(tax_rate or 0))
    total = Decimal(str(amount_incl_tax or 0))
    if rate <= 0:
        return float(total.quantize(Decimal('0.00'), rounding=ROUND_HALF_UP)), 0.0
    amount_excl_tax = (total / (1 + rate)).quantize(Decimal('0.00'), rounding=ROUND_HALF_UP)
    tax_amount = (total - amount_excl_tax).quantize(Decimal('0.00'), rounding=ROUND_HALF_UP)
    return float(amount_excl_tax), float(tax_amount)


FEE_TYPE_WHITELIST = {'入廊费', '日常维护费', '运维费', '维护费'}
PAYMENT_STATUS_WHITELIST = {'未开始', '待收缴', '部分收缴', '已收缴', '已逾期'}


def validate_record_values(fee_type: str, payment_status: str, amount_incl_tax: float,
                           tax_rate: float, actual_received_amount: float):
    """表单值白名单与范围校验，非法时抛 422 中文提示。"""
    if normalize_fee_type(fee_type) not in FEE_TYPE_WHITELIST:
        raise HTTPException(status_code=422, detail=f'费用类型非法：{fee_type}')
    if (payment_status or '').strip() not in PAYMENT_STATUS_WHITELIST:
        raise HTTPException(status_code=422, detail=f'收缴状态非法：{payment_status}')
    if amount_incl_tax < 0 or actual_received_amount < 0:
        raise HTTPException(status_code=422, detail='金额不能为负数')
    if not 0 <= tax_rate < 1:
        raise HTTPException(status_code=422, detail='税率应在 0~1 之间（如 9% 填 0.09）')
    # 状态与金额交叉校验
    status = (payment_status or '').strip()
    if status == '已收缴' and actual_received_amount < amount_incl_tax - 0.005:
        raise HTTPException(status_code=422, detail='状态为"已收缴"时，实收金额不能小于应收金额')
    if status == '部分收缴':
        if actual_received_amount <= 0:
            raise HTTPException(status_code=422, detail='状态为"部分收缴"时，实收金额必须大于 0')
        if actual_received_amount >= amount_incl_tax - 0.005:
            raise HTTPException(status_code=422, detail='状态为"部分收缴"时，实收金额应小于应收金额（已收清请改"已收缴"）')
    if status == '待收缴' and actual_received_amount > 0:
        raise HTTPException(status_code=422, detail='状态为"待收缴"时实收金额应为 0（有实收请改"部分收缴"或"已收缴"）')


def validate_record_relations(db: Session, company_id: int, pipeline_entry_id: int | None):
    company = db.query(Company).filter(Company.id == company_id).first()
    if not company:
        raise HTTPException(status_code=404, detail='单位不存在')

    if pipeline_entry_id is None:
        return

    pipeline_entry = db.query(PipelineEntry).filter(PipelineEntry.id == pipeline_entry_id).first()
    if not pipeline_entry:
        raise HTTPException(status_code=404, detail='入廊记录不存在')
    if pipeline_entry.company_id != company_id:
        raise HTTPException(status_code=422, detail='收费记录关联的入廊记录与所选单位不一致')


@router.get('/', response_class=HTMLResponse)
def list_records(
    request: Request,
    db: Session = Depends(get_db),
    company_id: str = '',
    status: str = '',
    project: str = '',
    notice: str = '',
    keyword: str = '',
    year: str = '',
    sort: str = '',
):
    from sqlalchemy import extract, or_
    from sqlalchemy.orm import joinedload
    selected_company_id = parse_optional_int(company_id)
    selected_project = (project or '').strip()
    selected_keyword = (keyword or '').strip()
    selected_year = parse_optional_int(year)
    query = db.query(FeeRecord).options(joinedload(FeeRecord.payments))
    if selected_company_id:
        query = query.filter(FeeRecord.company_id == selected_company_id)
    # '__overdue__' 是动态逾期筛选，不对应数据库里的状态值，查询后再在内存中过滤
    if status and status != OVERDUE_FILTER_VALUE:
        query = query.filter(FeeRecord.payment_status == status)
    if selected_project:
        query = query.filter(FeeRecord.project_name == selected_project)
    if selected_year:
        query = query.filter(extract('year', FeeRecord.planned_receivable_date) == selected_year)
    if selected_keyword:
        like = f'%{selected_keyword}%'
        query = query.join(Company, FeeRecord.company_id == Company.id).filter(or_(
            FeeRecord.project_name.ilike(like),
            FeeRecord.remark.ilike(like),
            Company.company_name.ilike(like),
            Company.short_name.ilike(like),
        ))
    sort_orders = {
        'due_asc': FeeRecord.planned_receivable_date.asc(),
        'amount_desc': FeeRecord.amount_incl_tax.desc(),
        'amount_asc': FeeRecord.amount_incl_tax.asc(),
    }
    records = query.order_by(sort_orders.get(sort, FeeRecord.planned_receivable_date.desc())).all()
    available_years = [row[0] for row in db.query(
        extract('year', FeeRecord.planned_receivable_date))
        .filter(FeeRecord.planned_receivable_date.isnot(None))
        .distinct().order_by(extract('year', FeeRecord.planned_receivable_date).desc()).all()
        if row[0] is not None]
    today = date.today()
    annotate_overdue(records, today)
    if status == OVERDUE_FILTER_VALUE:
        records = [record for record in records if record.is_overdue_view]
    overdue_count = sum(1 for record in records if record.is_overdue_view)
    summary_groups, summary_total = build_filtered_summary(records)
    companies = db.query(Company).order_by(Company.company_name).all()
    project_query = db.query(FeeRecord.project_name).filter(
        FeeRecord.project_name.isnot(None), FeeRecord.project_name != '')
    if selected_company_id:
        project_query = project_query.filter(FeeRecord.company_id == selected_company_id)
    project_names = [row[0] for row in project_query.distinct().order_by(FeeRecord.project_name).all()]
    reminder_settings = get_reminder_settings(db)
    last_run_at = get_setting(db, 'reminder.last_run_at', '')
    last_run_summary = get_setting(db, 'reminder.last_run_summary', '')

    return templates.TemplateResponse(
        'fee_records/list.html',
        {
            'request': request,
            'records': records,
            'companies': companies,
            'selected_company_id': selected_company_id,
            'selected_status': status,
            'project_names': project_names,
            'selected_project': selected_project,
            'selected_keyword': selected_keyword,
            'selected_year': selected_year,
            'available_years': available_years,
            'selected_sort': sort,
            'today_str': date.today().isoformat(),
            'overdue_filter_value': OVERDUE_FILTER_VALUE,
            'overdue_count': overdue_count,
            'summary_groups': summary_groups,
            'summary_total': summary_total,
            'notice': notice,
            'reminder_settings': reminder_settings,
            'last_run_at': last_run_at,
            'last_run_summary': last_run_summary,
            'title': '收费记录'
        }
    )


@router.post('/reminders/digest')
def run_digest_now(db: Session = Depends(get_db)):
    """手动触发一次每日汇总推送（忽略时间与去重检查）。"""
    message = run_daily_digest(db, force=True)
    return RedirectResponse(url=f'/fee-records/?notice={quote(message)}', status_code=303)


@router.post('/{record_id}/receive')
def receive_payment(record_id: int, request: Request, db: Session = Depends(get_db),
                    amount: float = Form(...), received_date: str = Form(''),
                    remark: str = Form('')):
    """登记一笔到账：写入流水并回写实收缓存与状态。"""
    from app.services.payment_service import add_payment
    record = db.query(FeeRecord).filter(FeeRecord.id == record_id).first()
    if not record:
        raise HTTPException(status_code=404, detail='收费记录不存在')
    if amount is None or amount <= 0:
        raise HTTPException(status_code=422, detail='收款金额必须大于 0')
    add_payment(db, record, amount, parse_date(received_date), remark)
    message = f'已登记收款 ¥{amount:,.2f}（{record.project_name or ""}），当前状态：{record.payment_status}。'
    return RedirectResponse(url=f'/fee-records/?notice={quote(message)}', status_code=303)


@router.post('/payments/{payment_id}/delete')
def delete_payment_route(payment_id: int, db: Session = Depends(get_db)):
    """删除一笔收款流水并回滚实收缓存（误登纠正）。"""
    from app.models.fee_payment import FeePayment
    from app.services.payment_service import delete_payment
    payment = db.query(FeePayment).filter(FeePayment.id == payment_id).first()
    if not payment:
        raise HTTPException(status_code=404, detail='收款流水不存在')
    record = delete_payment(db, payment)
    message = '已删除一笔收款流水，实收合计已回滚。'
    target = f'/fee-records/{record.id}/edit' if record else '/fee-records/'
    return RedirectResponse(url=f'{target}?notice={quote(message)}', status_code=303)


@router.get('/export')
def export_records(
    request: Request,
    db: Session = Depends(get_db),
    company_id: str = '',
    status: str = '',
    project: str = '',
):
    """导出收费记录为 Excel，沿用列表页的筛选条件。"""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    selected_company_id = parse_optional_int(company_id)
    selected_project = (project or '').strip()
    query = db.query(FeeRecord)
    if selected_company_id:
        query = query.filter(FeeRecord.company_id == selected_company_id)
    if status and status != OVERDUE_FILTER_VALUE:
        query = query.filter(FeeRecord.payment_status == status)
    if selected_project:
        query = query.filter(FeeRecord.project_name == selected_project)

    records = query.order_by(FeeRecord.planned_receivable_date.desc()).all()
    today = date.today()
    annotate_overdue(records, today)
    if status == OVERDUE_FILTER_VALUE:
        records = [record for record in records if record.is_overdue_view]

    wb = Workbook()
    ws = wb.active
    ws.title = '收费记录'
    headers = ['单位', '项目名称', '收费类型', '收费周期', '不含税金额', '税率', '税金',
               '含税金额(应收)', '应收时间', '实收金额', '实收时间', '收缴状态', '是否逾期', '备注']
    header_fill = PatternFill('solid', fgColor='D9E1F2')
    header_font = Font(bold=True)
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
    widths = [16, 24, 12, 14, 14, 8, 12, 14, 12, 12, 12, 10, 10, 24]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    for row_idx, record in enumerate(records, start=2):
        company_name = record.company.short_name or record.company.company_name if record.company else ''
        ws.cell(row=row_idx, column=1, value=safe_excel_cell(company_name))
        ws.cell(row=row_idx, column=2, value=safe_excel_cell(record.project_name or ''))
        ws.cell(row=row_idx, column=3, value=safe_excel_cell(record.fee_type or ''))
        ws.cell(row=row_idx, column=4, value=safe_excel_cell(record.charge_period or ''))
        ws.cell(row=row_idx, column=5, value=float(record.amount_excl_tax or 0))
        ws.cell(row=row_idx, column=6, value=float(record.tax_rate or 0))
        ws.cell(row=row_idx, column=7, value=float(record.tax_amount or 0))
        ws.cell(row=row_idx, column=8, value=float(record.amount_incl_tax or 0))
        ws.cell(row=row_idx, column=9, value=record.planned_receivable_date.isoformat() if record.planned_receivable_date else '')
        ws.cell(row=row_idx, column=10, value=float(record.actual_received_amount or 0))
        ws.cell(row=row_idx, column=11, value=record.actual_received_date.isoformat() if record.actual_received_date else '')
        ws.cell(row=row_idx, column=12, value=safe_excel_cell(record.payment_status or ''))
        ws.cell(row=row_idx, column=13, value='是' if record.is_overdue_view else '否')
        ws.cell(row=row_idx, column=14, value=safe_excel_cell(record.remark or ''))

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = quote('收费记录.xlsx')
    return StreamingResponse(
        buf,
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f"attachment; filename*=UTF-8''{filename}"},
    )


@router.post('/reminders/run')
def run_reminders():
    result = run_fee_reminders()
    return RedirectResponse(url=f'/fee-records/?notice={quote(result.message)}', status_code=303)


@router.get('/statement', response_class=HTMLResponse)
def company_statement(request: Request, company_id: int, db: Session = Depends(get_db)):
    """按单位生成对账单（打印页）：该单位全部收费记录与应收/已收/未收合计。"""
    company = db.query(Company).filter(Company.id == company_id).first()
    if not company:
        raise HTTPException(status_code=404, detail='单位不存在')
    records = (db.query(FeeRecord)
               .filter(FeeRecord.company_id == company_id)
               .order_by(FeeRecord.planned_receivable_date.asc(), FeeRecord.id.asc())
               .all())
    total_receivable = sum(float(r.amount_incl_tax or 0) for r in records)
    total_received = sum(float(r.actual_received_amount or 0) for r in records)
    return templates.TemplateResponse('fee_records/statement.html', {
        'request': request,
        'company': company,
        'records': records,
        'total_receivable': total_receivable,
        'total_received': total_received,
        'total_outstanding': total_receivable - total_received,
        'print_date': date.today().isoformat(),
        'title': '对账单',
    })


@router.get('/{record_id}/print', response_class=HTMLResponse)
def print_record_fee_sheet(record_id: int, request: Request, db: Session = Depends(get_db)):
    """按收费记录打印费用单：金额锚定记录本身（季度/首期折算都对得上账）。

    与按项目打印（恒年度全额）的区别：本单只出该记录对应期次的金额，
    明细行按「记录不含税额 ÷ 年度原始不含税额」等比缩放，尾行吸收分差，
    折后实际段直接取记录的含税/不含税/税金三值，与台账分毫不差。
    """
    from decimal import Decimal, ROUND_HALF_UP

    from app.routes.pipeline_entries import (
        PRINT_REGION_TITLE, PRINT_TARGET_ROWS, _fmt_money, _fmt_pct, _fmt_qty,
        calc_detail_amount, get_discount_or_default, infer_detail_unit)

    def q2d(v: Decimal) -> Decimal:
        return v.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

    record = db.query(FeeRecord).filter(FeeRecord.id == record_id).first()
    if not record or not record.pipeline_entry:
        return RedirectResponse(url='/fee-records/', status_code=303)
    entry = record.pipeline_entry
    is_entry = record.fee_type == '入廊费'
    fee_label = '入廊费' if is_entry else '运维费'
    tax_rate = float(record.tax_rate or 0)
    discount = get_discount_or_default(
        entry.entry_fee_discount if is_entry else entry.maintenance_fee_discount)

    record_incl = Decimal(str(record.amount_incl_tax or 0))
    record_excl = Decimal(str(record.amount_excl_tax or 0))
    record_tax = (Decimal(str(record.tax_amount)) if record.tax_amount is not None
                  else record_incl - record_excl)
    # 折前段由记录反推（折扣=1 时与记录完全一致）
    if discount and float(discount) != 1.0:
        orig_incl = q2d(record_incl / Decimal(str(discount)))
    else:
        orig_incl = record_incl
    if tax_rate > 0:
        orig_excl = q2d(orig_incl / (1 + Decimal(str(tax_rate))))
    else:
        orig_excl = orig_incl
    orig_tax = orig_incl - orig_excl

    # 明细行：原始明细不含税额等比缩放至折前不含税合计，尾行吸收分差
    bases = []
    for item in entry.details:
        qty = item.engineering_quantity or 0
        if is_entry:
            unit_price = item.entry_unit_price_excl_tax or 0
            amount = item.entry_amount_excl_tax or 0
        else:
            unit_price = item.maintenance_unit_price_excl_tax or 0
            amount = item.maintenance_amount_excl_tax or 0
        if not amount and qty and unit_price:
            amount = calc_detail_amount(qty, unit_price)
        bases.append((item, qty, unit_price, Decimal(str(amount or 0))))
    annual_orig_excl = sum((b[3] for b in bases), Decimal('0'))
    factor = (orig_excl / annual_orig_excl) if annual_orig_excl > 0 else Decimal('0')
    rows = []
    running = Decimal('0')
    for idx, (item, qty, unit_price, base) in enumerate(bases):
        if idx == len(bases) - 1:
            row_amount = orig_excl - running
        else:
            row_amount = q2d(base * factor)
            running += row_amount
        rows.append({
            'pipeline_type': item.pipeline_type or '',
            'specification': item.specification or '',
            'unit': infer_detail_unit(item.pipeline_type),
            'quantity': _fmt_qty(qty),
            'unit_price': _fmt_money(unit_price),
            'amount': _fmt_money(row_amount),
            'remark': item.remark or '',
        })

    cabin = entry.cabin_type or ''
    note_text = (
        f'根据管线入廊合同约定，本项目{cabin}{fee_label}按照上述金额的{_fmt_pct(discount)}收取。'
        '实际综合管廊有偿使用费如下：')
    if not is_entry and record.charge_period:
        note_text += f'本费用单对应收费期次：{record.charge_period}。'
    sheet_title = f'管廊{cabin}有偿使用费（{entry.project_name or ""}-{fee_label}'
    if not is_entry and record.charge_period:
        sheet_title += f'·{record.charge_period}'
    sheet_title += '）'
    fixed_rows = 1 + len(rows) + 3 + 4
    blank_rows = list(range(max(0, PRINT_TARGET_ROWS - fixed_rows)))
    return templates.TemplateResponse(
        'pipeline_entries/print.html',
        {
            'request': request,
            'region_title': PRINT_REGION_TITLE,
            'sheet_title': sheet_title,
            'qty_header': '入廊工程量' if is_entry else '运维工程量',
            'price_header': ('不含税入廊费（元/米·孔）' if is_entry
                             else '不含税运维费（元/米·年）'),
            'amount_header': '不含税入廊费合计' if is_entry else '不含税运维费合计',
            'rows': rows,
            'subtotal': _fmt_money(orig_excl),
            'tax_rate_text': _fmt_pct(tax_rate),
            'tax_amount': _fmt_money(orig_tax),
            'total_incl_tax': _fmt_money(orig_incl),
            'note_text': note_text,
            'actual_subtotal': _fmt_money(record_excl),
            'actual_tax': _fmt_money(record_tax),
            'actual_total': _fmt_money(record_incl),
            'blank_rows': blank_rows,
            'title': f'打印{fee_label} - {entry.project_name or ""}',
        },
    )


@router.get('/new', response_class=HTMLResponse)
def new_record(request: Request, db: Session = Depends(get_db)):
    companies = db.query(Company).order_by(Company.company_name).all()
    pipeline_entries = db.query(PipelineEntry).order_by(PipelineEntry.id.desc()).all()
    entry_fee_map = {}
    for entry in pipeline_entries:
        entry_actual_fee, maintenance_actual_fee = calculate_entry_actual_fees(entry)
        entry_fee_map[entry.id] = {
            'entry_actual_fee': entry_actual_fee,
            'maintenance_actual_fee': maintenance_actual_fee,
        }
    statuses = ['未开始', '待收缴', '部分收缴', '已收缴', '已逾期']

    return templates.TemplateResponse(
        'fee_records/form.html',
        {
            'request': request,
            'record': None,
            'companies': companies,
            'pipeline_entries': pipeline_entries,
            'entry_fee_map': entry_fee_map,
            'statuses': statuses,
            'standard_tax_rates_json': safe_json_dumps(get_all_standard_tax_rates(db)),
            'title': '新增收费记录'
        }
    )


@router.post('/new')
def create_record(
    company_id: int = Form(...),
    pipeline_entry_id: str | None = Form(None),
    project_name: str = Form(''),
    fee_type: str = Form(...),
    charge_period: str = Form(''),
    period_year: str | None = Form(None),
    period_quarter: str | None = Form(None),
    amount_incl_tax: str | None = Form(None),
    tax_rate: str | None = Form(None),
    planned_receivable_date: str = Form(''),
    remind_date: str = Form(''),
    latest_payment_date: str = Form(''),
    actual_received_amount: str | None = Form(None),
    actual_received_date: str = Form(''),
    payment_status: str = Form('待收缴'),
    is_invoiced: str = Form('否'),
    remark: str = Form(''),
    db: Session = Depends(get_db),
):
    # 空字符串归一化（浏览器提交空数字框/空下拉为 ""）
    pipeline_entry_id = parse_optional_int(pipeline_entry_id)
    period_year = parse_optional_int(period_year)
    period_quarter = parse_optional_int(period_quarter)
    amount_incl_tax = parse_float_or(amount_incl_tax, 0.0)
    tax_rate = parse_float_or(tax_rate, 0.0)
    actual_received_amount = parse_float_or(actual_received_amount, 0.0)
    validate_record_relations(db, company_id, pipeline_entry_id)
    validate_record_values(fee_type, payment_status, amount_incl_tax, tax_rate, actual_received_amount)
    
    # 动态获取税率：启用中的收费标准（最新生效），找不到则用前端传入值兜底
    actual_tax_rate = get_standard_tax_rate(db, fee_type)
    if actual_tax_rate is None:
        actual_tax_rate = float(tax_rate)
    if not (0 <= float(actual_tax_rate) < 1):
        raise HTTPException(status_code=422, detail='税率异常（%s），请检查收费标准或表单税率' % actual_tax_rate)

    amount_excl_tax, tax_amount = calc_excl_tax_from_incl_tax(amount_incl_tax, actual_tax_rate)

    new_record = FeeRecord(
            company_id=company_id,
            pipeline_entry_id=pipeline_entry_id,
            project_name=project_name,
            fee_type=fee_type,
            charge_period=charge_period,
            period_year=period_year,
            period_quarter=period_quarter,
            amount_excl_tax=amount_excl_tax,
            tax_rate=actual_tax_rate,  # 使用动态计算出的税率
            tax_amount=tax_amount,
            amount_incl_tax=amount_incl_tax,
            planned_receivable_date=parse_date(planned_receivable_date),
            remind_date=parse_date(remind_date),
            latest_payment_date=parse_date(latest_payment_date),
            actual_received_amount=actual_received_amount,
            actual_received_date=parse_date(actual_received_date),
            payment_status=payment_status,
            is_invoiced=is_invoiced,
            remark=remark,
        )
    db.add(new_record)
    db.flush()
    if actual_received_amount and actual_received_amount > 0:
        from app.models.fee_payment import FeePayment
        db.add(FeePayment(fee_record_id=new_record.id,
                          amount=actual_received_amount,
                          received_date=parse_date(actual_received_date),
                          remark='建记录时登记'))
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=422, detail='收费记录保存失败，请确认关联数据有效') from exc
    return RedirectResponse(url='/fee-records/', status_code=303)


@router.get('/{record_id}/edit', response_class=HTMLResponse)
def edit_record(record_id: int, request: Request, db: Session = Depends(get_db)):
    record = db.query(FeeRecord).filter(FeeRecord.id == record_id).first()
    if not record:
        return RedirectResponse(url='/fee-records/', status_code=303)

    companies = db.query(Company).order_by(Company.company_name).all()
    pipeline_entries = db.query(PipelineEntry).order_by(PipelineEntry.id.desc()).all()
    entry_fee_map = {}
    for entry in pipeline_entries:
        entry_actual_fee, maintenance_actual_fee = calculate_entry_actual_fees(entry)
        entry_fee_map[entry.id] = {
            'entry_actual_fee': entry_actual_fee,
            'maintenance_actual_fee': maintenance_actual_fee,
        }
    statuses = ['未开始', '待收缴', '部分收缴', '已收缴', '已逾期']

    return templates.TemplateResponse(
        'fee_records/form.html',
        {
            'request': request,
            'record': record,
            'companies': companies,
            'pipeline_entries': pipeline_entries,
            'entry_fee_map': entry_fee_map,
            'statuses': statuses,
            'standard_tax_rates_json': safe_json_dumps(get_all_standard_tax_rates(db)),
            'title': '编辑收费记录'
        }
    )


@router.post('/{record_id}/edit')
def update_record(
    record_id: int,
    company_id: int = Form(...),
    pipeline_entry_id: str | None = Form(None),
    project_name: str = Form(''),
    fee_type: str = Form(...),
    charge_period: str = Form(''),
    period_year: str | None = Form(None),
    period_quarter: str | None = Form(None),
    amount_incl_tax: str | None = Form(None),
    tax_rate: str | None = Form(None),
    planned_receivable_date: str = Form(''),
    remind_date: str = Form(''),
    latest_payment_date: str = Form(''),
    actual_received_amount: str | None = Form(None),
    actual_received_date: str = Form(''),
    payment_status: str = Form('待收缴'),
    is_invoiced: str = Form('否'),
    remark: str = Form(''),
    db: Session = Depends(get_db),
):
    record = db.query(FeeRecord).filter(FeeRecord.id == record_id).first()
    if not record:
        return RedirectResponse(url='/fee-records/', status_code=303)

    # 表单空字符串安全解析，避免 422
    pipeline_entry_id = parse_optional_int(pipeline_entry_id)
    period_year = parse_optional_int(period_year)
    period_quarter = parse_optional_int(period_quarter)
    amount_incl_tax = parse_float_or(amount_incl_tax, 0.0)
    submitted_tax_rate = parse_float_or(tax_rate, None if tax_rate is None else 0.0)
    actual_received_amount = parse_float_or(actual_received_amount, 0.0)

    validate_record_relations(db, company_id, pipeline_entry_id)
    validate_record_values(fee_type, payment_status, amount_incl_tax,
                           submitted_tax_rate if submitted_tax_rate is not None else 0.0,
                           actual_received_amount)

    # 编辑时保留历史税率：仅费用类型变更或用户显式修改税率时才重取
    old_tax_rate = float(record.tax_rate or 0)
    if fee_type != record.fee_type:
        actual_tax_rate = get_standard_tax_rate(db, fee_type)
        if actual_tax_rate is None:
            actual_tax_rate = submitted_tax_rate if submitted_tax_rate is not None else old_tax_rate
    elif submitted_tax_rate is not None and abs(submitted_tax_rate - old_tax_rate) > 1e-9:
        actual_tax_rate = submitted_tax_rate
    else:
        actual_tax_rate = old_tax_rate
    if not (0 <= float(actual_tax_rate) < 1):
        raise HTTPException(status_code=422, detail='税率异常（%s），请检查收费标准或表单税率' % actual_tax_rate)

    amount_excl_tax, tax_amount = calc_excl_tax_from_incl_tax(amount_incl_tax, actual_tax_rate)

    record.company_id = company_id
    record.pipeline_entry_id = pipeline_entry_id
    record.project_name = project_name
    record.fee_type = fee_type
    record.charge_period = charge_period
    record.period_year = period_year
    record.period_quarter = period_quarter
    record.amount_excl_tax = amount_excl_tax
    record.tax_rate = actual_tax_rate  # 更新为动态计算的税率
    record.tax_amount = tax_amount
    record.amount_incl_tax = amount_incl_tax
    record.planned_receivable_date = parse_date(planned_receivable_date)
    record.remind_date = parse_date(remind_date)
    record.latest_payment_date = parse_date(latest_payment_date)
    record.actual_received_amount = actual_received_amount
    record.actual_received_date = parse_date(actual_received_date)
    record.payment_status = payment_status
    record.is_invoiced = is_invoiced
    record.remark = remark

    # 已有收款流水的记录：实收总额/日期/状态以流水汇总为准，防止表单手改与流水脱节
    if record.payments:
        from app.services.payment_service import sync_record_from_payments
        sync_record_from_payments(db, record)

    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=422, detail='收费记录更新失败，请确认关联数据有效') from exc
    return RedirectResponse(url='/fee-records/', status_code=303)


@router.post('/{record_id}/delete')
def delete_record(record_id: int, db: Session = Depends(get_db)):
    record = db.query(FeeRecord).filter(FeeRecord.id == record_id).first()
    if not record:
        return RedirectResponse(url='/fee-records/', status_code=303)

    db.delete(record)
    db.commit()
    return RedirectResponse(url='/fee-records/', status_code=303)
